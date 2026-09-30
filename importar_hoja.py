#!/usr/bin/env python3
"""Sube a Notion lo que ya tienes en la hoja "Finanzas Personales 2025".

1. En Google Drive abre la hoja, Archivo > Descargar > Microsoft Excel (.xlsx).
2. python3 importar_hoja.py "Finanzas Personales 2025.xlsx" --probar   (muestra que haria)
3. python3 importar_hoja.py "Finanzas Personales 2025.xlsx"             (sube)

Lee las pestañas Gastos, Ingresos y Ahorro (las metas). Las demas no se tocan.
Cada fila subida queda anotada en data/importados.json: si lo corres de nuevo
no duplica nada.
"""
import argparse
import hashlib
import json
import sys
import time
from datetime import date, datetime

import categorias as C
from config import DATA, IMPORTADOS
from notion import METAS, MOVIMIENTOS, Notion, NotionError, cargar_bases, p_date, p_number, p_select, p_title

# categorias de la hoja que no se llaman igual aca
EQUIVALENCIAS = {
    "citas medicas": "Salud", "medicinas": "Salud", "salud": "Salud",
    "gastos financieros": "Gastos financieros", "comida y restaurantes": "Comida y restaurantes",
    "supermercado": "Alimentación", "alimentacion": "Alimentación", "mercado": "Alimentación",
    "servicios del hogar": "Servicios del hogar", "servicios": "Servicios del hogar",
    "matricula": "Educación", "educacion": "Educación", "cursos": "Educación",
    "ropa": "Ropa y cuidado personal", "cuidado personal": "Ropa y cuidado personal",
    "taxi": "Transporte", "movilidad": "Transporte", "transporte": "Transporte",
    "freshtunes": "Música y distribución", "agregadoras": "Música y distribución",
}
PRIORIDADES = {"muy importante": "Muy importante", "importante": "Importante", "no importante": "Puede esperar"}


def fecha(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = str(v or "").strip().split(" ")[0]
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    return None


def numero(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).replace("S/", "").replace("$", "").replace(" ", "").strip()
    if "," in t and "." in t:
        t = t.replace(",", "") if t.rfind(".") > t.rfind(",") else t.replace(".", "").replace(",", ".")
    else:
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def filas(hoja) -> list:
    """Las filas como diccionarios usando la primera fila con 'Fecha' como encabezado."""
    datos = list(hoja.iter_rows(values_only=True))
    for i, fila in enumerate(datos[:10]):
        nombres = [str(x or "").strip() for x in fila]
        if "Fecha" in nombres:
            return [dict(zip(nombres, f)) for f in datos[i + 1:]]
    return []


def categoria(tipo: str, nombre: str, desc: str) -> str:
    n = C.normal(nombre).strip()
    if n in EQUIVALENCIAS and EQUIVALENCIAS[n] in C.nombres(tipo):
        return EQUIVALENCIAS[n]
    for cat in C.nombres(tipo):
        if C.normal(cat) == n:
            return cat
    return C.buscar_categoria(tipo, nombre) or C.adivinar(tipo, desc) or C.otros(tipo)


def movimientos(libro) -> list:
    out = []
    for pestaña, es_ingreso in (("Gastos", False), ("Ingresos", True)):
        if pestaña not in libro.sheetnames:
            print("  (no hay pestaña %s)" % pestaña)
            continue
        for k, f in enumerate(filas(libro[pestaña])):
            d = fecha(f.get("Fecha"))
            imp = numero(f.get("Importe"))
            mn = numero(f.get("Importe en MN") if "Importe en MN" in f else f.get("Importe MN"))
            if not d or not (imp or mn):
                continue
            imp = imp or mn
            mn = mn if mn is not None else imp
            desc = str(f.get("Descripción") or "").strip()
            cat_hoja = str(f.get("Categoría") or "").strip()
            cuenta = C.normal(str(f.get("Cuenta") or ""))
            if es_ingreso:
                tipo = "Ingreso"
            elif C.normal(cat_hoja) == "ahorro":
                tipo = "Ahorro"
            elif "inversion" in cuenta:
                tipo = "Inversión"
            else:
                tipo = "Gasto"
            if tipo == "Inversión":
                cat = "Inmueble" if C.normal(cat_hoja).startswith("vivienda") else (C.adivinar("Inversión", desc) or "Otras inversiones")
            else:
                cat = categoria(tipo, cat_hoja, desc)
            moneda = str(f.get("Moneda") or "PEN").strip().upper() or "PEN"
            if moneda not in ("PEN", "USD", "EUR"):
                moneda, imp = "PEN", mn
            medio = C.buscar_medio(str(f.get("Medio de Pago") or ""))
            clave = hashlib.sha1(("%s|%d|%s|%s|%s" % (pestaña, k, d, desc, imp)).encode()).hexdigest()[:16]
            out.append({"clave": clave, "tipo": tipo, "categoria": cat, "cat_hoja": cat_hoja, "desc": desc or cat_hoja or tipo,
                        "monto": imp, "moneda": moneda, "soles": mn, "medio": medio, "fecha": d})
    return out


def metas(libro) -> list:
    if "Ahorro" not in libro.sheetnames:
        return []
    datos = list(libro["Ahorro"].iter_rows(values_only=True))
    if not datos:
        return []
    enc = [str(x or "").strip() for x in datos[0]]
    out = []
    for f in datos[1:]:
        r = dict(zip(enc, f))
        nombre = str(r.get("Item") or "").strip()
        obj = numero(r.get("Objetivo"))
        if nombre and obj:
            out.append({"meta": nombre, "objetivo": obj, "ahorrado": numero(r.get("Ahorrado")) or 0,
                        "prioridad": PRIORIDADES.get(C.normal(str(r.get("Relevancia") or "")).strip())})
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("archivo")
    ap.add_argument("--probar", action="store_true", help="muestra lo que haria sin subir nada")
    a = ap.parse_args(argv)
    try:
        import openpyxl
    except ImportError:
        print("Falta openpyxl: .venv/bin/pip install -r requirements.txt")
        return 1
    libro = openpyxl.load_workbook(a.archivo, data_only=True, read_only=True)
    movs = movimientos(libro)
    lista_metas = metas(libro)
    try:
        hechos = set(json.loads(IMPORTADOS.read_text()))
    except (OSError, ValueError):
        hechos = set()
    nuevos = [m for m in movs if m["clave"] not in hechos]

    resumen = {}
    for m in nuevos:
        k = (m["tipo"], m["categoria"])
        resumen[k] = resumen.get(k, 0) + m["soles"]
    print("Movimientos en la hoja: %d · ya subidos: %d · por subir: %d" % (len(movs), len(movs) - len(nuevos), len(nuevos)))
    for (tipo, cat), v in sorted(resumen.items(), key=lambda kv: -kv[1]):
        print("  %-10s %-26s S/ %12s" % (tipo, cat, "{:,.2f}".format(v)))
    sin_cat = sorted({m["cat_hoja"] for m in nuevos if m["categoria"] == C.otros(m["tipo"]) and m["cat_hoja"] and m["tipo"] != "Ahorro"})
    if sin_cat:
        print("Categorías de la hoja que quedan en 'Otros': %s" % ", ".join(sin_cat))
    if lista_metas:
        print("Metas: %s" % ", ".join("%s (%s de %s)" % (m["meta"], m["ahorrado"], m["objetivo"]) for m in lista_metas))
    if a.probar:
        print("\n--probar: no subí nada.")
        return 0

    bases = cargar_bases()
    if MOVIMIENTOS not in bases:
        print("Falta data/bases.json. Corre primero: python3 setup_notion.py")
        return 1
    n = Notion()
    DATA.mkdir(parents=True, exist_ok=True)
    subidos = 0
    try:
        for m in nuevos:
            tc = round(m["soles"] / m["monto"], 4) if m["monto"] else 1
            n.crear_pagina(bases[MOVIMIENTOS], {
                "Descripción": p_title(m["desc"]), "Tipo": p_select(m["tipo"]), "Categoría": p_select(m["categoria"]),
                "Grupo": p_select(C.grupo(m["categoria"]) if m["tipo"] == "Gasto" else "—"),
                "Monto": p_number(m["monto"]), "Moneda": p_select(m["moneda"]), "Tipo de cambio": p_number(tc),
                "Monto S/": p_number(m["soles"]), "Medio de pago": p_select(m["medio"]), "Fecha": p_date(m["fecha"]),
                "Origen": p_select("Hoja 2025"),
            })
            hechos.add(m["clave"])
            subidos += 1
            if subidos % 25 == 0:
                IMPORTADOS.write_text(json.dumps(sorted(hechos)))
                print("  %d/%d" % (subidos, len(nuevos)), flush=True)
            time.sleep(0.35)   # Notion acepta unas 3 escrituras por segundo
        if lista_metas and METAS in bases:
            ya = {C.normal(f.get("Meta") or "") for f in n.consultar(bases[METAS], limite=200)}
            for m in lista_metas:
                if C.normal(m["meta"]) in ya:
                    continue
                props = {"Meta": p_title(m["meta"]), "Objetivo S/": p_number(m["objetivo"]), "Ahorrado S/": p_number(m["ahorrado"])}
                if m["prioridad"]:
                    props["Prioridad"] = p_select(m["prioridad"])
                n.crear_pagina(bases[METAS], props)
                print("  meta %s" % m["meta"])
    except NotionError as exc:
        print("Notion cortó después de %d filas: %s\nVuelve a correr el comando: sigue donde quedó." % (subidos, exc))
        return 1
    finally:
        IMPORTADOS.write_text(json.dumps(sorted(hechos)))
    print("Listo: %d movimientos subidos." % subidos)
    return 0


if __name__ == "__main__":
    sys.exit(main())
