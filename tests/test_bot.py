import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import finanzas as F  # noqa: E402
from bot import Bot, codigo_salida, leer_patrimonio  # noqa: E402
from telegram import TelegramError  # noqa: E402
from tests.fakes import BASES, FakeNotion, FakeTelegram  # noqa: E402

YO = 7


def mensaje(texto, user=YO, chat=YO, tipo="private"):
    return {"update_id": 1, "message": {"chat": {"id": chat, "type": tipo}, "from": {"id": user}, "text": texto}}


def boton(data, user=YO, chat=YO):
    return {"update_id": 2, "callback_query": {"id": "cb", "from": {"id": user}, "data": data,
                                               "message": {"chat": {"id": chat, "type": "private"}, "message_id": 9}}}


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self._aj = F.AJUSTES
        F.AJUSTES = Path(self._tmp.name) / "ajustes.json"
        self.tg, self.n = FakeTelegram(), FakeNotion()
        self.bot = Bot(self.tg, self.n, BASES, {YO})

    def tearDown(self):
        F.AJUSTES = self._aj
        self._tmp.cleanup()

    def di(self, texto):
        self.bot.procesar(mensaje(texto))
        return self.tg.ultimo

    def toca(self, etiqueta):
        self.bot.procesar(boton(self.tg.data_de(etiqueta)))
        return self.tg.ultimo

    @property
    def movs(self):
        return self.n.dbs["db-mov"]


class Autorizacion(Base):
    def test_sin_usuarios_dice_el_id(self):
        b = Bot(self.tg, self.n, BASES, set())
        b.procesar(mensaje("45 almuerzo", user=555, chat=555))
        self.assertIn("555", self.tg.ultimo)
        self.assertEqual(self.movs, [])

    def test_extrano_una_sola_vez(self):
        self.bot.procesar(mensaje("hola", user=9, chat=9))
        self.bot.procesar(mensaje("hola", user=9, chat=9))
        self.assertEqual(len(self.tg.enviados), 1)

    def test_grupos_se_ignoran(self):
        self.bot.procesar(mensaje("45 almuerzo", tipo="group"))
        self.assertEqual(self.tg.enviados, [])


class Anotar(Base):
    def test_gasto_queda_en_notion(self):
        t = self.di("45 almuerzo interbank")
        self.assertIn("Comida y restaurantes", t)
        self.assertIn("S/ 45.00", t)
        f = self.movs[0]
        self.assertEqual((f["Tipo"], f["Monto S/"], f["Medio de pago"], f["Grupo"], f["Origen"]),
                         ("Gasto", 45, "Interbank", "Deseo", "Telegram"))

    def test_categoria_desconocida_pregunta(self):
        self.di("80 cosas raras")
        self.toca("Movilidad")
        self.assertEqual(self.movs[0]["Categoría"], "Movilidad")
        self.assertEqual(self.movs[0]["Grupo"], "Necesidad")

    def test_cambiar_categoria(self):
        self.di("30 uber")
        self.toca("Cambiar categoría")
        self.toca("Viajes")
        self.assertEqual(self.movs[0]["Categoría"], "Viajes")

    def test_deshacer_con_boton_y_comando(self):
        self.di("30 uber")
        self.toca("Deshacer")
        self.assertEqual(self.movs, [])
        self.di("20 taxi")
        self.di("/deshacer")
        self.assertEqual(self.movs, [])
        self.assertIn("No hay nada", self.di("/deshacer"))

    def test_no_entendi(self):
        self.assertIn("No encontré el monto", self.di("almuerzo"))
        self.assertEqual(self.movs, [])

    def test_comando_forzado(self):
        self.di("/ingreso 200 regalo")
        self.assertEqual(self.movs[0]["Tipo"], "Ingreso")

    def test_dolares_muestra_conversion(self):
        self.di("/tc 3.70")
        self.di("+100 usd facebook")
        t = self.tg.enviados[-2][1]
        self.assertIn("USD 100.00 = S/ 370.00", t)


class AhorroYMetas(Base):
    def test_ahorro_suma_a_la_meta_y_se_deshace(self):
        self.di("/meta Fondo de emergencia 20000")
        t = self.di("ahorro 500 emergencia")
        self.assertIn("S/ 500.00 de S/ 20,000.00", t)
        self.assertEqual(self.n.dbs["db-met"][0]["Ahorrado S/"], 500)
        self.toca("Deshacer")
        self.assertEqual(self.n.dbs["db-met"][0]["Ahorrado S/"], 0)

    def test_metas_lista(self):
        self.di("/meta Auto 100k")
        self.assertIn("S/ 100,000.00", self.di("/metas"))


class Presupuesto(Base):
    def test_fijar_y_avisar(self):
        self.assertIn("S/ 100.00 · $ 26.67 · ₽ 2,222 al mes", self.di("/presupuesto comida 100"))
        self.assertIn("🟡", self.di("85 almuerzo"))
        self.assertIn("Pasaste el presupuesto", self.di("20 cafe"))
        t = self.di("/presupuesto")
        self.assertIn("🔴 Comida y restaurantes", t)

    def test_categoria_invalida(self):
        self.assertIn("No reconozco", self.di("/presupuesto zzz 100"))


class Patrimonio(Base):
    def test_activos_y_deudas(self):
        self.di("/activo Interbank 5200")
        self.di("/deuda Tarjeta Ripley 1200")
        t = self.di("/activo Interbank 6000")
        self.assertIn("antes S/ 5,200.00", t)
        self.assertIn("Patrimonio neto: <b>S/ 4,800.00 · $ 1,280.00 · ₽ 106,667</b>", t)
        self.assertIn("Lo que debes", self.di("/patrimonio"))

    def test_leer_patrimonio(self):
        self.assertEqual(leer_patrimonio("Activo", "Binance 800 usd"), ("Binance", 800, "USD", "Cripto"))
        self.assertEqual(leer_patrimonio("Activo", "Auto 45000"), ("Auto", 45000, "PEN", "Vehículo"))
        self.assertEqual(leer_patrimonio("Pasivo", "Tarjeta Ripley 1,200"), ("Tarjeta Ripley", 1200, "PEN", "Tarjeta de crédito"))


class Deudas(Base):
    @property
    def deudas(self):
        return self.n.dbs["db-deu"]

    def test_deuda_va_a_la_base_deudas(self):
        t = self.di("/deuda Tarjeta Ripley 1200")
        self.assertIn("Total de deudas: <b>S/ 1,200.00 · $ 320.00 · ₽ 26,667</b>", t)
        d = self.deudas[0]
        self.assertEqual((d["Deuda"], d["Tipo"], d["Saldo"], d["Monto original"], d["Estado"]),
                         ("Tarjeta Ripley", "Tarjeta de crédito", 1200, 1200, "Activa"))
        self.assertEqual(self.n.dbs["db-pat"], [])
        t = self.di("/deuda Tarjeta Ripley 1500")
        self.assertIn("antes S/ 1,200.00", t)
        self.assertEqual(len(self.deudas), 1)

    def test_cuenta_en_el_patrimonio_neto(self):
        self.di("/activo Interbank 5200")
        self.di("/deuda Préstamo BCP 2000")
        t = self.di("/patrimonio")
        self.assertIn("Patrimonio neto: S/ 3,200.00", t)
        self.assertIn("Préstamo BCP", t)

    def test_pago_baja_el_saldo_hasta_pagarla(self):
        self.di("/deuda Tarjeta Ripley 1000")
        t = self.di("/pago ripley 300")
        self.assertIn("Te queda: <b>S/ 700.00 · $ ", t)
        self.assertEqual(self.deudas[0]["Saldo"], 700)
        self.assertEqual(self.movs, [])  # pagar una deuda no es un gasto nuevo
        t = self.di("/pago ripley 700")
        self.assertIn("pagada por completo", t)
        self.assertEqual(self.deudas[0]["Estado"], "Pagada")
        self.assertIn("No tienes deudas activas", self.di("/deudas"))

    def test_pago_en_otra_moneda(self):
        F.fijar_tipo_de_cambio("USD", 4.0)
        self.di("/deuda Juan 100 usd")
        self.assertEqual(self.deudas[0]["Tipo"], "Otra")
        self.di("/pago juan 200 soles")
        self.assertEqual(self.deudas[0]["Saldo"], 50)

    def test_lista_y_errores(self):
        self.di("/deuda Tarjeta Ripley 1200")
        self.di("/deuda Préstamo BCP 15000")
        self.n.dbs["db-deu"][1].update({"Tasa anual": 0.18, "Cuota mensual": 800, "Día de pago": 5})
        t = self.di("/deudas")
        self.assertIn("Deudas: S/ 16,200.00", t)
        self.assertLess(t.index("Préstamo BCP"), t.index("Tarjeta Ripley"))
        self.assertIn("tasa 18%", t)
        self.assertIn("Cuotas al mes: S/ 800.00", t)
        self.assertIn("No encuentro", self.di("/pago visa 10"))
        self.assertIn("¿A qué deuda le pagaste?", self.di("/pago"))
        self.assertIn("Pagar Préstamo BCP", [t for t, _ in self.tg.botones()][0])
        self.assertIn("Escribe el nombre", self.di("/deuda"))

    def test_avalancha_en_consejos(self):
        self.di("/deuda Tarjeta Ripley 1200")
        self.di("/deuda Préstamo BCP 15000")
        self.n.dbs["db-deu"][0]["Tasa anual"] = 0.65
        self.n.dbs["db-deu"][1]["Tasa anual"] = 0.18
        self.di("+3000 sueldo")
        self.assertIn("avalancha: paga el mínimo en todas y lo extra a Tarjeta Ripley", self.di("/consejos"))

    def test_sin_base_deudas_sigue_en_patrimonio(self):
        bases = {k: v for k, v in self.bot.bases.items() if k != "Deudas"}
        b = Bot(self.tg, self.n, bases, {YO})
        b.procesar(mensaje("/deuda Tarjeta Ripley 1200"))
        self.assertEqual(self.n.dbs["db-pat"][0]["Clase"], "Pasivo")

    def test_editar_en_notion(self):
        F.fijar_tipo_de_cambio("RUB", 0.05)
        self.di("/deuda Tarjeta Falabella 1000")
        self.di("/deuda Luis 2000 rub")
        fal, luis = self.n.dbs["db-deu"]
        fal["Saldo"] = 400                       # lo cambio a mano en Notion, sin tocar Saldo S/
        luis["Saldo"] = 0
        self.n.dbs["db-deu"].append({"_id": "%032d" % 999, "_orden": 2, "Deuda": "Janet", "Saldo": 300})
        t = self.di("/deudas")
        self.assertIn("Deudas: S/ 700.00", t)    # 400 + 300 (Janet, sin moneda = soles)
        self.assertNotIn("Luis", t)
        self.assertEqual((fal["Saldo S/"], luis["Estado"]), (400, "Pagada"))
        self.assertEqual((self.n.dbs["db-deu"][2]["Saldo S/"], self.n.dbs["db-deu"][2]["Estado"]), (300, "Activa"))

    def test_tres_monedas(self):
        F.fijar_tipo_de_cambio("USD", 3.38)
        F.fijar_tipo_de_cambio("RUB", 0.0428)
        self.assertEqual(F.s3(1000), "S/ 1,000.00 · $ 295.86 · ₽ 23,364")
        self.assertEqual(F.s3(-20), "-S/ 20.00 · -$ 5.92 · -₽ 467")
        self.di("+100 usd facebook")
        t = self.tg.enviados[-2][1]
        self.assertIn("USD 100.00 = S/ 338.00 · $ 100.00 · ₽ 7,897", t)
        self.assertEqual((self.movs[0]["Monto USD"], self.movs[0]["Monto RUB"]), (100, 7897))
        self.di("/deuda Luis 2941")
        self.di("/deudas")
        d = self.n.dbs["db-deu"][0]
        self.assertEqual((d["Saldo USD"], d["Saldo RUB"]), (870.12, 68715))

    def test_comando_tc(self):
        self.assertIn("de respaldo", self.di("/tc"))
        self.assertIn("fijo desde ahora", self.di("/tc 3.38"))
        t = self.di("/tc")
        self.assertIn("1 dólar = S/ 3.38 · <i>fijado por ti</i>", t)
        self.di("/tc rub 0.0428")
        self.assertIn("1 dólar = ₽ 78.97", self.di("/tc"))
        t = self.di("/tc auto")
        self.assertIn("vuelve a ser automático", t)
        self.assertNotIn("fijado por ti", t)

    def test_cuenta_en_dolares_sigue_al_cambio(self):
        F.fijar_tipo_de_cambio("USD", 3.38)
        self.di("/activo Interbank dólares 1497 usd")
        F.fijar_tipo_de_cambio("USD", 3.50)
        self.assertIn("S/ 5,239.50", self.di("/patrimonio"))
        self.assertEqual(self.n.dbs["db-pat"][0]["Valor S/"], 5239.5)

    def test_menu_tiene_deudas(self):
        self.di("/start")
        self.assertIn("No tienes deudas activas", self.toca("Deudas"))


class Resumenes(Base):
    def test_hoy_semana_mes(self):
        self.di("+3000 sueldo")
        self.di("45 almuerzo")
        self.assertIn("S/ 45.00", self.di("/hoy"))
        self.assertIn("Gastos: <b>S/ 45.00 · $ 12.00 · ₽ 1,000</b>", self.di("/semana"))
        self.assertIn("tasa de ahorro 98%", self.di("/mes"))
        self.bot.procesar(boton("m:presupuesto"))
        self.assertIn("Presupuesto", self.tg.ultimo)

    def test_consejos_y_metodos(self):
        self.di("+1000 sueldo")
        self.di("1200 alquiler")
        self.assertIn("🔴", self.di("/consejos"))
        self.assertIn("50/30/20", self.di("/metodos"))

    def test_ultimos(self):
        self.di("45 almuerzo")
        self.di("12 taxi")
        t = self.di("/ultimos")
        self.assertLess(t.index("Taxi"), t.index("Almuerzo"))


class Salida(unittest.TestCase):
    def test_codigos(self):
        self.assertEqual(codigo_salida(TelegramError("x", 401)), 2)
        self.assertEqual(codigo_salida(TelegramError("x", 409)), 3)
        self.assertEqual(codigo_salida(TelegramError("x", 500)), 1)


if __name__ == "__main__":
    unittest.main()


class Formularios(Base):
    def setUp(self):
        super().setUp()
        import formularios
        self._fa = formularios.AJUSTES
        formularios.AJUSTES = F.AJUSTES

    def tearDown(self):
        import formularios
        formularios.AJUSTES = self._fa
        super().tearDown()

    def test_gasto_completo_como_el_google_form(self):
        self.di("/gasto")
        self.assertIn("Nuevo gasto", self.tg.enviados[-2][1])
        self.assertIn("¿Qué fecha?", self.tg.ultimo)
        self.toca("Ayer")
        self.assertIn("cuenta", self.tg.ultimo)
        self.toca("Salud")
        self.toca("T-Bank")
        self.assertIn("¿Crédito o débito?", self.tg.ultimo)
        self.toca("Débito")
        self.toca("Medicina")
        self.toca("RUB")
        self.assertIn("descripción", self.tg.ultimo)
        self.di("pastillas para la gripe")
        self.assertIn("¿Cuánto fue?", self.tg.ultimo)
        self.assertIn("número", self.di("mucho"))
        t = self.di("1.250,50")
        self.assertIn("revisa antes de guardar", t)
        self.assertIn("RUB 1,250.50", t)
        self.assertEqual(self.movs, [])          # nada se guarda sin tocar Guardar
        self.toca("Guardar")
        f = self.movs[0]
        self.assertEqual((f["Tipo"], f["Cuenta"], f["Medio de pago"], f["Categoría"], f["Moneda"], f["Monto"], f["Descripción"]),
                         ("Gasto", "Salud", "T-Bank", "Medicina", "RUB", 1250.5, "Pastillas para la gripe"))
        self.assertEqual(f["Fecha"], (F.hoy() - F.timedelta(days=1)).isoformat())
        self.assertAlmostEqual(f["Monto S/"], round(1250.5 * F.tipo_de_cambio("RUB"), 2))
        self.assertIn("cuenta Salud", self.tg.ultimo)

    def test_corregir_antes_de_guardar(self):
        self.di("/gasto")
        for b in ("Hoy", "Gastos", "Interbank", "Débito", "Supermercado", "PEN"):
            self.toca(b)
        self.toca("Omitir")
        self.di("80")
        self.toca("Corregir")
        self.toca("Importe")
        t = self.di("85")
        self.assertIn("PEN 85.00", t)            # vuelve directo al resumen
        self.toca("Guardar")
        self.assertEqual((self.movs[0]["Monto"], self.movs[0]["Descripción"]), (85, "Supermercado"))

    def test_lo_ultimo_elegido_sale_primero(self):
        self.di("/gasto")
        self.toca("Hoy")
        self.toca("Gastos")
        self.toca("KuCoin")
        self.toca("Cancelar")
        self.di("/gasto")
        self.toca("Hoy")
        self.toca("Gastos")
        self.assertTrue(self.tg.botones()[0][0].startswith("KuCoin"))

    def test_ingreso_y_fecha_escrita(self):
        self.di("/ingreso")
        self.di("15/09")
        self.toca("PayPal")
        self.toca("Facebook")
        self.toca("USD")
        self.di("pago de septiembre")
        self.di("1943")
        self.toca("Guardar")
        f = self.movs[0]
        self.assertEqual((f["Tipo"], f["Categoría"], f["Medio de pago"], f["Moneda"], f["Fecha"][5:]),
                         ("Ingreso", "Facebook", "PayPal", "USD", "09-15"))

    def test_ahorro_a_una_meta(self):
        self.di("/meta Fondo de emergencia 20000")
        self.di("/ahorro")
        self.toca("Hoy")
        self.toca("Fondo de emergencia")
        self.toca("Interbank")
        self.toca("PEN")
        self.di("500")
        self.toca("Guardar")
        self.assertEqual(self.n.dbs["db-met"][0]["Ahorrado S/"], 500)
        self.assertEqual(self.movs[0]["Categoría"], "Fondo de emergencia")

    def test_cancelar_y_boton_viejo(self):
        self.di("/gasto")
        viejo = self.tg.data_de("Hoy")
        self.toca("Hoy")
        self.bot.procesar(boton(viejo))
        self.assertIn("pregunta anterior", self.tg.enviados[-2][1])
        self.di("/cancelar")
        self.assertIn("Cancelado", self.tg.ultimo)
        self.assertEqual(self.movs, [])

    def test_cuenta_inversion_no_es_gasto(self):
        self.di("/gasto")
        for b in ("Hoy", "Inversión", "Interbank", "Débito", "Vivienda", "PEN"):
            self.toca(b)
        self.di("autovaluo depa")
        self.di("2174")
        self.toca("Guardar")
        m = F.mes()
        r = F.resumir(F.movimientos(self.n, BASES, m.desde, m.hasta), m)
        self.assertEqual((r.gastos, r.inversion), (0, 2174))

    def _gasto(self, medio, tarjeta=None, monto="100", moneda="PEN"):
        self.di("/gasto")
        for b in ("Hoy", "Gastos", medio):
            self.toca(b)
        if tarjeta:
            self.toca(tarjeta)
        for b in ("Supermercado", moneda, "Omitir"):
            self.toca(b)
        self.di(monto)
        return self.toca("Guardar")

    def test_credito_suma_a_la_deuda_de_la_tarjeta(self):
        self.di("/deuda Banco Falabella 10900")
        t = self._gasto("Falabella", "Crédito")
        self.assertIn("se sumó a Banco Falabella. Ahora debes S/ 11,000.00", t)
        self.assertEqual(self.movs[0]["Tarjeta"], "Crédito")
        self.assertEqual(len(self.n.dbs["db-deu"]), 1)
        self.assertEqual(self.n.dbs["db-deu"][0]["Saldo"], 11000)
        self.di("/deshacer")
        self.assertEqual(self.n.dbs["db-deu"][0]["Saldo"], 10900)
        self.assertEqual(self.movs, [])

    def test_credito_crea_la_deuda_si_no_existe(self):
        t = self._gasto("BBVA", "Crédito", "250")
        self.assertIn("creé la deuda Tarjeta BBVA", t)
        d = self.n.dbs["db-deu"][0]
        self.assertEqual((d["Deuda"], d["Tipo"], d["Saldo"], d["Estado"], d["Acreedor"]),
                         ("Tarjeta BBVA", "Tarjeta de crédito", 250, "Activa", "BBVA"))
        self._gasto("BBVA", "Crédito", "50")
        self.assertEqual((len(self.n.dbs["db-deu"]), self.n.dbs["db-deu"][0]["Saldo"]), (1, 300))

    def test_credito_en_otra_moneda_y_tarjeta_pagada(self):
        F.fijar_tipo_de_cambio("RUB", 0.05)
        self.di("/deuda Tarjeta T-Bank 1000 rub")
        self.di("/pago t-bank 1000")
        self.assertEqual(self.n.dbs["db-deu"][0]["Estado"], "Pagada")
        self._gasto("T-Bank", "Crédito", "10", "PEN")        # S/ 10 = 200 RUB
        d = self.n.dbs["db-deu"][0]
        self.assertEqual((d["Saldo"], d["Estado"]), (200, "Activa"))

    def test_debito_y_efectivo_no_tocan_deudas(self):
        self._gasto("Interbank", "Débito")
        self.assertEqual(self.movs[0]["Tarjeta"], "Débito")
        self.di("/gasto")
        for b in ("Hoy", "Gastos", "Efectivo"):
            self.toca(b)
        self.assertIn("¿Qué categoría?", self.tg.ultimo)    # efectivo no pregunta credito o debito
        self.assertEqual(self.n.dbs["db-deu"], [])

    def test_corregir_medio_a_tarjeta_pregunta_credito(self):
        self.di("/gasto")
        for b in ("Hoy", "Gastos", "Efectivo", "Supermercado", "PEN", "Omitir"):
            self.toca(b)
        self.di("40")
        self.toca("Corregir")
        self.assertNotIn("Tarjeta", [t for t, _ in self.tg.botones()])
        self.toca("Medio de pago")
        self.toca("SIP")
        self.assertIn("¿Crédito o débito?", self.tg.ultimo)
        t = self.toca("Crédito")
        self.assertIn("revisa antes de guardar", t)
        self.assertIn("Tarjeta: <b>Crédito</b>", t)
        self.toca("Guardar")
        self.assertEqual(self.n.dbs["db-deu"][0]["Deuda"], "Tarjeta SIP")

    def test_rapido_con_la_palabra_credito(self):
        t = self.di("45 almuerzo cmr credito")
        self.assertIn("Banco Falabella", t)
        self.assertEqual(self.movs[0]["Descripción"], "Almuerzo")
        self.di("30 taxi credito")                          # sin tarjeta: no hay a que deuda sumarlo
        self.assertEqual(len(self.n.dbs["db-deu"]), 1)

    def test_medios_nuevos(self):
        self.di("/gasto")
        self.toca("Hoy")
        self.toca("Gastos")
        nombres = [t for t, _ in self.tg.botones()]
        for m in ("T-Bank", "Falabella", "SIP", "KuCoin"):
            self.assertIn(m, nombres)
        for m in ("CMR", "Tinkoff", "Ripley", "Scotiabank", "Paxful"):
            self.assertNotIn(m, nombres)
        self.toca("KuCoin")
        self.assertIn("¿Qué categoría?", self.tg.ultimo)    # KuCoin no es tarjeta

    def test_menu_tiene_los_formularios(self):
        self.di("/start")
        self.bot.procesar(boton(self.tg.data_de("Gasto")))
        self.assertIn("¿Qué fecha?", self.tg.ultimo)


class CuentasYAhorro(Base):
    """El ingreso sube la cuenta donde entra, el gasto (no a credito) la baja, y al cobrar se ofrece ahorrar."""

    def cuenta(self, nombre):
        return next(f for f in self.n.dbs["db-pat"] if f["Nombre"] == nombre)

    def test_ingreso_y_gasto_mueven_la_cuenta(self):
        F.fijar_tipo_de_cambio("USD", 3.38)
        self.di("/activo Interbank dólares 1497 usd")
        self.di("/activo Interbank soles 463")
        self.di("+100 usd facebook interbank")
        self.assertIn("Interbank dólares ahora tiene S/ 5,397.86 · $ 1,597.00", self.tg.enviados[-2][1])
        self.assertEqual(self.cuenta("Interbank dólares")["Valor"], 1597)
        self.assertEqual(self.cuenta("Interbank soles")["Valor"], 463)      # la de soles no se toca
        t = self.di("63 almuerzo interbank")
        self.assertIn("Interbank soles ahora tiene S/ 400.00", t)
        self.di("/deshacer")
        self.assertEqual(self.cuenta("Interbank soles")["Valor"], 463)

    def test_credito_y_sin_cuenta_no_mueven(self):
        self.di("/activo Falabella 500")
        self.di("50 almuerzo falabella credito")
        self.assertEqual(self.cuenta("Falabella")["Valor"], 500)
        t = self.di("50 almuerzo bbva")
        self.assertIn("/activo BBVA 1000 pen", t)
        self.assertEqual(len(self.n.dbs["db-pat"]), 1)

    def test_retirar_efectivo_pasa_del_banco_al_efectivo(self):
        F.fijar_tipo_de_cambio("USD", 3.38)
        self.di("/activo Interbank soles 463")
        self.di("/activo Interbank dólares 1497 usd")
        self.di("/start")
        self.toca("Retirar efectivo")
        self.toca("Interbank soles")
        self.di("200")
        self.assertEqual(self.cuenta("Interbank soles")["Valor"], 263)
        self.assertEqual(self.cuenta("Efectivo")["Valor"], 200)
        self.assertEqual(self.cuenta("Efectivo")["Tipo"], "Efectivo y bancos")
        self.assertEqual(len(self.movs), 0)                                  # no es un gasto
        self.di("/retirar")
        self.toca("Interbank soles")
        self.di("50")
        self.assertEqual(self.cuenta("Efectivo")["Valor"], 250)
        self.di("/retirar")
        self.assertFalse(any(t.startswith("🏦 Efectivo") for t, _ in self.tg.botones()))   # no se retira del efectivo
        self.toca("Interbank dólares")
        self.di("100")
        self.assertEqual(self.cuenta("Interbank dólares")["Valor"], 1397)
        self.assertEqual(self.cuenta("Efectivo dólares")["Valor"], 100)
        self.toca("Deshacer retiro")
        self.assertEqual(self.cuenta("Interbank dólares")["Valor"], 1497)
        self.assertEqual(self.cuenta("Efectivo dólares")["Valor"], 0)
        t = self.di("30 almuerzo efectivo")
        self.assertIn("Efectivo ahora tiene S/ 220.00", t)

    def test_plin_y_yape_salen_de_interbank(self):
        self.di("/activo Interbank soles 463")
        self.di("/activo BCP 300")
        self.assertIn("Interbank soles ahora tiene S/ 443.00", self.di("20 taxi plin"))
        self.assertIn("Interbank soles ahora tiene S/ 428.00", self.di("15 menu yape"))
        self.assertEqual(self.movs[-1]["Medio de pago"], "Yape")                # el medio queda como Yape
        self.assertEqual(self.cuenta("BCP")["Valor"], 300)                      # BCP no se toca

    def test_yape_que_salio_de_bcp(self):
        self.di("/activo Interbank soles 463")
        self.di("/activo BCP 300")
        self.di("15 menu yape")
        self.assertEqual(self.cuenta("Interbank soles")["Valor"], 448)
        t = self.toca("Salió de BCP")
        self.assertIn("Interbank soles vuelve a S/ 463.00", t)
        self.assertEqual(self.cuenta("Interbank soles")["Valor"], 463)
        self.assertEqual(self.cuenta("BCP")["Valor"], 285)
        self.di("/deshacer")                                                   # deshacer devuelve a BCP
        self.assertEqual(self.cuenta("BCP")["Valor"], 300)
        self.di("20 taxi plin")
        self.assertFalse(any("Salió de" in t for t, _ in self.tg.botones()))  # Plin siempre es Interbank

    def test_retirar_sin_cuentas_y_cancelar(self):
        self.assertIn("/activo", self.di("/retirar"))
        self.di("/activo BCP 100")
        self.di("/retirar")
        self.toca("Cancelar")
        self.assertEqual(self.cuenta("BCP")["Valor"], 100)

    def test_pagate_primero_20_por_ciento_a_una_meta(self):
        self.di("/meta Emergencia 10000")
        self.di("+3000 sueldo")
        self.assertIn("¿Separas algo para ahorro?", self.tg.ultimo)
        self.toca("20%")
        self.toca("Emergencia")
        ahorro = self.movs[-1]
        self.assertEqual((ahorro["Tipo"], ahorro["Monto"], ahorro["Categoría"]), ("Ahorro", 600, "Fondo de emergencia"))
        self.assertEqual(self.n.dbs["db-met"][0]["Ahorrado S/"], 600)

    def test_otro_monto_y_ahorro_general(self):
        self.di("+3000 sueldo")
        self.toca("Otro monto")
        self.assertIn("número", self.di("mucho"))
        self.di("250")
        self.toca("Ahorro general")
        self.assertEqual((self.movs[-1]["Monto"], self.movs[-1]["Categoría"]), (250, "Ahorro general"))
        self.assertEqual(len(self.movs), 2)
        self.di("45 almuerzo")                 # despues de elegir, lo escrito vuelve a ser un gasto
        self.assertEqual(self.movs[-1]["Tipo"], "Gasto")

    def test_no_esta_vez(self):
        self.di("+3000 sueldo")
        self.toca("No esta vez")
        self.assertEqual(len(self.movs), 1)
        self.di("+500 sueldo")
        self.di("/cancelar")
        self.assertIn("No separo nada", self.tg.ultimo)


class Excel(Base):
    def test_excel_con_graficos(self):
        from io import BytesIO
        from openpyxl import load_workbook
        F.fijar_tipo_de_cambio("USD", 3.38)
        F.fijar_tipo_de_cambio("RUB", 0.0428)
        self.di("/deuda Banco SIP 16224")
        self.di("/deuda Tarjeta T-Bank 21756 rub")
        self.di("/activo Interbank dólares 1497 usd")
        self.di("/meta Fondo de emergencia 17240")
        self.di("/presupuesto supermercado 855")
        self.di("+3000 sueldo")
        self.di("120 supermercado")
        self.di("45 almuerzo")
        self.di("/excel")
        chat, nombre, contenido, texto = self.tg.documentos[-1]
        self.assertTrue(nombre.startswith("Finanzas ") and nombre.endswith(".xlsx"))
        wb = load_workbook(BytesIO(contenido))
        self.assertEqual(wb.sheetnames, ["Resumen", "Deudas", "Cuentas", "Metas", "Presupuesto", "Por mes", "Categorías", "Movimientos"])
        self.assertEqual(wb["Deudas"]["A2"].value, "Banco SIP")
        self.assertEqual(wb["Deudas"]["G3"].value, 21756)                 # T-Bank en rublos
        self.assertEqual(wb["Cuentas"]["F2"].value, F.en_rublos(1497 * 3.38))
        self.assertEqual(wb["Movimientos"].max_row, 4)                    # cabecera + 3 movimientos
        resumen = {wb["Resumen"].cell(row=k, column=1).value: wb["Resumen"].cell(row=k, column=2).value for k in range(5, 17)}
        self.assertEqual(resumen["📊 Patrimonio neto"], round(1497 * 3.38 - 16224 - 21756 * 0.0428, 2))
        for hoja in ("Resumen", "Deudas", "Cuentas", "Metas", "Presupuesto", "Por mes", "Categorías"):
            self.assertTrue(wb[hoja]._charts, hoja)

    def test_excel_muestra_el_limite(self):
        from io import BytesIO
        from openpyxl import load_workbook
        F.fijar_tipo_de_cambio("RUB", 0.05)
        self.di("/limite 1500 rub")
        self.di("1200 rub supermercado")
        self.di("/excel")
        ws = load_workbook(BytesIO(self.tg.documentos[-1][2]))["Resumen"]
        filas = {ws.cell(row=k, column=1).value: [ws.cell(row=k, column=c).value for c in range(2, 11)] for k in range(1, 30)}
        self.assertEqual(filas["Hoy"][2], 1500)        # puedes ₽
        self.assertEqual(filas["Hoy"][5], 1200)        # gastado ₽
        self.assertEqual(filas["Esta semana"][2], 10500)

    def test_excel_vacio_no_falla(self):
        self.di("/excel")
        self.assertEqual(len(self.tg.documentos), 1)

    def test_boton_excel(self):
        self.di("/start")
        self.toca("Excel")
        self.assertEqual(len(self.tg.documentos), 1)


class PagarConBotones(Base):
    def deuda(self, nombre):
        return next(f for f in self.n.dbs["db-deu"] if f["Deuda"] == nombre)

    def cuenta(self, nombre):
        return next(f for f in self.n.dbs["db-pat"] if f["Nombre"] == nombre)

    def test_cuota_y_cuenta_de_donde_salio(self):
        self.di("/deuda Banco Falabella 15751")
        self.deuda("Banco Falabella")["Cuota mensual"] = 1374
        self.di("/activo Interbank soles 2000")
        self.di("/deudas")
        self.toca("Pagar Banco Falabella")
        self.assertIn("¿Cuánto pagaste?", self.tg.ultimo)
        self.toca("Cuota")
        self.assertIn("Te queda: <b>S/ 14,377.00", self.tg.enviados[-2][1])
        self.assertIn("¿De qué cuenta salió el pago?", self.tg.ultimo)
        self.toca("Interbank soles")
        self.assertIn("Interbank soles ahora tiene S/ 626.00", self.tg.ultimo)
        self.assertEqual((self.deuda("Banco Falabella")["Saldo"], self.cuenta("Interbank soles")["Valor"]), (14377, 626))
        self.assertEqual(self.movs, [])                       # pagar una deuda no es un gasto nuevo

    def test_otro_monto_desde_cuenta_en_dolares(self):
        F.fijar_tipo_de_cambio("USD", 3.38)
        self.di("/deuda Multa por no ir a votar 200")
        self.di("/activo Interbank dólares 1497 usd")
        self.di("/pago")
        self.toca("Pagar Multa")
        self.toca("Otro monto")
        self.assertIn("número", self.di("mucho"))
        self.di("200")
        self.assertIn("pagada por completo", self.tg.enviados[-2][1])
        self.toca("Interbank dólares")
        self.assertEqual(self.cuenta("Interbank dólares")["Valor"], round(1497 - 200 / 3.38, 2))

    def test_todo_sin_descontar_y_cancelar(self):
        self.di("/deuda Luis 2941")
        self.di("/activo SIP 51")
        self.di("/deudas")
        self.toca("Pagar Luis")
        self.toca("Cancelar")
        self.assertEqual(self.deuda("Luis")["Saldo"], 2941)
        self.di("/deudas")
        self.toca("Pagar Luis")
        self.toca("Todo")
        self.toca("No descontar")
        self.assertEqual((self.deuda("Luis")["Estado"], self.cuenta("SIP")["Valor"]), ("Pagada", 51))

    def test_escrito_tambien_pregunta_la_cuenta(self):
        self.di("/deuda Janet 3800")
        self.di("/activo Interbank soles 463")
        self.di("/pago janet 300")
        self.assertIn("¿De qué cuenta salió el pago?", self.tg.ultimo)
        self.di("45 almuerzo")                                 # seguir anotando no se rompe
        self.assertEqual(self.movs[-1]["Tipo"], "Gasto")


class LimiteDiaADia(Base):
    def setUp(self):
        super().setUp()
        F.fijar_tipo_de_cambio("RUB", 0.05)

    def test_sin_limite_explica(self):
        self.assertIn("Aún no tienes límite", self.di("/limite"))
        self.assertNotIn("📏", self.di("45 almuerzo"))

    def test_hoy_semana_mes(self):
        F.fijar_tipo_de_cambio("USD", 4.0)
        self.assertIn("S/ 75.00 · $ 18.75 · ₽ 1,500 por día", self.di("/limite 1500 rub"))
        t = self.di("1200 rub supermercado")
        self.assertIn("🟡 Quedan hoy: S/ 15.00 · $ 3.75 · ₽ 300", t)      # 80%: amarillo
        self.assertIn("🟢 Quedan en la semana: S/ 465.00 · $ 116.25 · ₽ 9,300", t)
        dias_mes = F.mes().dias
        self.assertIn("Quedan en el mes: S/ {:,.2f}".format((1500 * dias_mes - 1200) * 0.05), t)
        t = self.di("55000 rub alquiler departamento")       # Vivienda: gasto fijo, no cuenta
        self.assertEqual(self.movs[-1]["Categoría"], "Vivienda")
        self.assertNotIn("Quedan hoy", t)
        t = self.di("500 rub taxi")
        self.assertIn("🔴 Te pasaste hoy: S/ 10.00 · $ 2.50 · ₽ 200", t)
        h = self.di("/hoy")
        self.assertIn("Día a día", h)
        self.assertIn("Gastado: S/ 85.00 · $ 21.25 · ₽ 1,700", h)
        self.assertIn("Te pasaste: S/ 10.00", h)
        sem = self.di("/semana")
        self.assertIn("Límite: S/ 525.00 · $ 131.25 · ₽ 10,500", sem)

    def test_inicio_muestra_el_limite_arriba(self):
        t = self.di("/start")
        self.assertIn("/limite 1500 rub", t)
        self.assertNotIn("como hablas", t)
        self.di("/limite 1500 rub")
        self.di("1200 rub supermercado")
        self.di("/deuda Banco SIP 16224")
        self.di("/meta Pasajes 10140")
        t = self.di("/start")
        primeras = t.split("\n")[:6]
        self.assertEqual(primeras[0], "📏 Quedan hoy: S/ 15.00 · $ 4.00 · ₽ 300")   # lo primero que se ve
        self.assertIn("Finanzas", primeras[1])
        self.assertIn("Para gastar en el día a día", primeras[3])
        self.assertIn("<b>Hoy</b>", primeras[4])
        self.assertIn("Gastado: S/ 60.00", primeras[5])
        self.assertIn("Deudas: <b>S/ 16,224.00", t)
        self.assertIn("Pasajes 0%", t)
        self.assertTrue(any("Ajustar" in x for x, _ in self.tg.botones()))   # el menu sigue abajo
        self.assertIn("/gasto", self.di("/ayuda"))

    def test_start_no_fija_nada(self):
        self.di("/limite 1500 rub")
        self.di("/start")
        self.di("100 rub cafe")
        self.assertFalse(hasattr(self.tg, "fijados"))
        self.assertIn("Quedan hoy", self.tg.enviados[-2][1] if "Quedan hoy" not in self.tg.ultimo else self.tg.ultimo)

    def test_quitar(self):
        self.di("/limite 64")
        self.assertIn("S/ 64.00 · $", self.di("/limite"))
        self.assertIn("Quité el límite", self.di("/limite 0"))
        self.assertIn("Aún no tienes límite", self.di("/limite"))


class PresupuestoAnualYFuera(Base):
    def setUp(self):
        super().setUp()
        import formularios
        self._fa = formularios.AJUSTES
        formularios.AJUSTES = F.AJUSTES

    def tearDown(self):
        import formularios
        formularios.AJUSTES = self._fa
        super().tearDown()

    def test_mensual_y_anual_separados(self):
        self.di("/presupuesto suscripciones 100")
        self.assertIn("al año", self.di("/presupuesto suscripciones anual 600"))
        pre = self.n.dbs["db-pre"][0]
        self.assertEqual((pre["Mensual S/"], pre["Anual S/"]), (100, 600))
        self.di("450 icloud anual")
        t = self.tg.enviados[-2][1]                              # el ultimo pregunta si es suscripcion nueva
        self.assertEqual((self.movs[-1]["Categoría"], self.movs[-1]["Frecuencia"]), ("Suscripciones", "Anual"))
        self.assertEqual(self.movs[-1]["Descripción"], "Icloud")
        self.assertIn("Pagos anuales de Suscripciones: S/ 450.00 de S/ 600.00", t)
        t = self.di("68 claude pro suscripcion")                # mensual: no se mezcla con el anual
        self.assertNotIn("Pasaste", t)
        p = self.di("/presupuesto")
        self.assertIn("S/ 68.00 de S/ 100.00", p)
        self.assertIn("Pagos anuales de", p)
        self.assertIn("S/ 450.00 de S/ 600.00", p)
        self.di("200 vpn anual suscripcion")
        self.assertIn("🔴 Pasaste el presupuesto anual", self.tg.enviados[-2][1])

    def test_formulario_pregunta_mensual_o_anual(self):
        self.di("/presupuesto vivienda 100")
        self.di("/presupuesto vivienda anual 600")
        self.di("/gasto")
        for b in ("Hoy", "Gastos", "Efectivo", "Vivienda"):
            self.toca(b)
        self.assertIn("pago anual", self.tg.ultimo)
        self.toca("Anual")
        for b in ("PEN", "Omitir"):
            self.toca(b)
        self.di("120")
        self.assertIn("Pago: <b>Anual</b>", self.tg.ultimo)
        self.toca("Guardar")
        self.assertEqual(self.movs[-1]["Frecuencia"], "Anual")
        self.di("/gasto")
        for b in ("Hoy", "Gastos", "Efectivo", "Supermercado"):   # sin tope anual: no pregunta
            self.toca(b)
        self.assertIn("¿En qué moneda?", self.tg.ultimo)

    def test_fuera_del_presupuesto(self):
        self.di("/presupuesto supermercado 800")
        t = self.di("150 consulta medica")
        self.assertEqual(self.movs[-1]["Categoría"], "Citas médicas")
        self.assertIn("⚠️ Fuera del presupuesto: Citas médicas", t)
        self.assertNotIn("Fuera del presupuesto", self.di("100 supermercado"))
        self.di("/limite 64")
        self.assertNotIn("Fuera del presupuesto", self.di("30 almuerzo"))    # dia a dia: lo controla el limite
        self.assertIn("⚠️ <b>Fuera del presupuesto</b>: Citas médicas S/ 150.00 · total S/ 150.00", self.di("/mes"))
        self.assertIn("fuera del presupuesto", self.di("/presupuesto"))

    def test_sin_presupuesto_no_marca_nada(self):
        self.assertNotIn("Fuera del presupuesto", self.di("150 consulta medica"))


class Suscripciones(Base):
    def sus(self, nombre):
        return next(f for f in self.n.dbs["db-sus"] if f["Suscripción"] == nombre)

    def test_nueva_mensual_y_despues_la_reconoce(self):
        F.fijar_tipo_de_cambio("USD", 3.38)
        self.di("/presupuesto suscripciones 500")
        self.di("20 usd claude pro")
        self.assertIn("es una suscripción nueva", self.tg.ultimo)
        t = self.toca("Mensual")
        self.assertIn("Agregué <b>Claude pro</b> (mensual)", t)
        self.assertIn("S/ 67.60 al mes + S/ 0.00 al año", t)
        self.assertIn("🟢 Máximo: S/ 500.00 al mes · quedan S/ 432.40", t)
        s = self.sus("Claude pro")
        self.assertEqual((s["Cada (meses)"], s["Estado"]), (1, "Activa"))
        self.assertEqual(s["Próximo pago"], F.sumar_meses(F.hoy(), 1).isoformat())
        n = len(self.tg.enviados)
        self.di("20 usd claude")                       # la segunda vez no pregunta
        self.assertIn("🔁 Claude pro · mensual · próximo pago", self.tg.ultimo)
        self.assertNotIn("suscripción nueva", " ".join(t for _, t, _ in self.tg.enviados[n:]))
        self.assertEqual(len(self.n.dbs["db-sus"]), 1)

    def test_anual_va_al_tope_anual_y_no_infla_el_mes(self):
        self.di("/presupuesto suscripciones 500")
        self.di("150 terabox")
        self.toca("Anual")
        self.assertEqual(self.movs[-1]["Frecuencia"], "Anual")
        self.assertEqual(self.n.dbs["db-pre"][0]["Anual S/"], 150)
        self.assertEqual(self.n.dbs["db-pre"][0]["Mensual S/"], 500)   # el maximo no se toca
        self.assertIn("S/ 0.00 de S/ 500.00", self.di("/presupuesto"))

    def test_cada_n_meses_y_sin_nombre(self):
        self.di("/gasto")
        for b in ("Hoy", "Gastos", "Efectivo", "Suscripciones", "PEN", "Omitir"):
            self.toca(b)
        self.di("90")
        self.toca("Guardar")
        self.toca("Otro (meses)")
        self.assertIn("número", self.di("muchos"))
        self.di("3")
        self.assertIn("¿Cómo se llama", self.tg.ultimo)
        self.assertIn("cada 3 meses", self.di("Gimnasio app"))
        s = self.sus("Gimnasio app")
        self.assertEqual((s["Cada (meses)"], s["Por mes S/"]), (3, 30))
        self.assertEqual(self.n.dbs["db-pre"][0]["Anual S/"], 360)

    def test_pasarse_del_maximo_pago_unico_y_cancelar(self):
        self.di("/presupuesto suscripciones 50")
        self.di("80 netflix")
        self.assertIn("🔴 Máximo: S/ 50.00 al mes · te pasas por S/ 30.00", self.toca("Mensual"))
        self.di("40 adobe")
        self.toca("Pago único")
        self.assertEqual(len(self.n.dbs["db-sus"]), 1)
        self.assertIn("Netflix", self.di("/suscripciones"))
        self.assertIn("Cancelé Netflix", self.di("/suscripcion cancelar netflix"))
        self.assertEqual(self.sus("Netflix")["Estado"], "Cancelada")
        self.assertIn("Todavía no hay", self.di("/suscripciones"))

    def test_aviso_de_renovacion_en_el_resumen_diario(self):
        self.di("20 claude")
        self.toca("Mensual")
        self.sus("Claude")["Próximo pago"] = (F.hoy() + F.timedelta(days=2)).isoformat()
        self.assertIn("🔁 Claude se renueva el", self.di("/hoy"))


class VerYAjustar(Base):
    def test_ver_muestra_las_opciones(self):
        self.di("/start")                                     # todo esta en la pantalla principal
        nombres = [t for t, _ in self.tg.botones()]
        for x in ("Límite", "Presupuesto", "Deudas", "Patrimonio", "Metas", "Suscripciones", "Excel", "Consejos"):
            self.assertTrue(any(x in n for n in nombres), x)
        self.assertIn("Límite del día a día", self.toca("Límite"))

    def test_ajustar_pide_el_dato_y_lo_guarda(self):
        F.fijar_tipo_de_cambio("RUB", 0.05)
        self.di("/start")
        self.bot.procesar(boton("m:ajustar"))
        self.assertIn("¿Cuánto puedes gastar por día", self.toca("Límite diario"))
        self.assertIn("₽ 1,500 por día", self.di("1500 rub"))
        self.bot.procesar(boton("m:ajustar"))
        self.toca("Deuda nueva")
        self.di("Tarjeta BBVA 1200")
        self.assertEqual(self.n.dbs["db-deu"][0]["Deuda"], "Tarjeta BBVA")
        self.bot.procesar(boton("m:ajustar"))
        self.toca("Saldo de una cuenta")
        self.di("T-Bank 25000 rub")
        self.assertEqual((self.n.dbs["db-pat"][0]["Nombre"], self.n.dbs["db-pat"][0]["Moneda"]), ("T-Bank", "RUB"))
        self.di("45 almuerzo")                                # despues vuelve a anotar normal
        self.assertEqual(self.movs[-1]["Tipo"], "Gasto")

    def test_ajustar_cancelar_y_pagar(self):
        self.di("/deuda Luis 2941")
        self.di("/start")
        self.bot.procesar(boton("m:ajustar"))
        self.toca("Meta de ahorro")
        self.toca("Cancelar")
        self.di("45 almuerzo")
        self.assertEqual(self.movs[-1]["Tipo"], "Gasto")
        self.assertEqual(self.n.dbs["db-met"], [])
        self.di("/start")
        self.bot.procesar(boton("m:ajustar"))
        self.toca("Pagar una deuda")
        self.assertIn("¿A qué deuda le pagaste?", self.tg.ultimo)


class RecordatoriosYPlan(Base):
    def rec(self, nombre, dia, tipo="Otro", monto=None, moneda="PEN", categoria=None, deuda=""):
        from notion import p_number, p_select, p_text, p_title
        props = {"Recordatorio": p_title(nombre), "Día": p_number(dia), "Tipo": p_select(tipo), "Moneda": p_select(moneda),
                 "Estado": p_select("Activo"), "Deuda": p_text(deuda)}
        if monto:
            props["Monto"] = p_number(monto)
        if categoria:
            props["Categoría"] = p_select(categoria)
        return self.n.crear_pagina("db-rec", props)["id"]

    def test_avisos_del_dia_y_botones(self):
        import avisos
        hoy = F.hoy()
        self.rec("Retirar dinero de Payoneer", hoy.day, "Retiro")
        self.rec("Pagar servicios", hoy.day, "Pago", 350, "PEN", "Servicios")
        self.rec("Otro día", (hoy.day % 28) + 1, "Otro")
        msgs = avisos.mensajes(self.n, self.bot.bases)
        textos = [t for t, _ in msgs]
        self.assertTrue(any("Retirar dinero de Payoneer" in t for t in textos))
        self.assertFalse(any("Otro día" in t for t in textos))
        pago = next(b for t, b in msgs if "Pagar servicios" in t)
        self.bot.procesar(boton(pago[0][0][1]))                  # ✅ Pagado S/ 350
        self.assertEqual((self.movs[-1]["Categoría"], self.movs[-1]["Monto"]), ("Servicios", 350))
        retiro = next(b for t, b in msgs if "Payoneer" in t)
        self.bot.procesar(boton(retiro[0][1][1]))               # ⏰ Mañana
        self.assertIn("Te lo recuerdo mañana", self.tg.ultimo)
        textos = [t for t, _ in avisos.mensajes(self.n, self.bot.bases)]
        self.assertFalse(any("Pagar servicios" in t for t in textos))   # ya esta hecho este mes
        self.assertFalse(any("Payoneer" in t for t in textos))          # pospuesto para mañana

    def test_fin_de_mes_y_cuota_de_banco(self):
        self.di("/deuda Banco Falabella 15751")
        self.n.dbs["db-deu"][0]["Cuota mensual"] = 1374
        self.di("/activo Interbank soles 5000")
        self.rec("Cuota Banco Falabella", 31, "Deuda", 1374, "PEN", deuda="Banco Falabella")
        r = F.recordatorios(self.n, self.bot.bases)[0]
        fin = F.mes().hasta
        self.assertEqual(F.dia_del_mes(31, fin), fin.day)       # el 31 cae el ultimo dia del mes
        t = self.di("/pagos")
        self.assertIn("Cuota Banco Falabella", t)
        self.toca("Pagar")
        self.assertEqual(self.n.dbs["db-deu"][0]["Saldo"], 15751 - 1374)
        self.assertIn("¿De qué cuenta salió el pago?", self.tg.ultimo)
        self.assertIn("✅ Cuota Banco Falabella", self.di("/pagos"))

    def test_proxima_fecha_para_el_calendario(self):
        hoy = F.hoy()
        self.rec("Retirar dinero de PayPal", 23, "Retiro")
        F.recordatorios(self.n, self.bot.bases)
        fila = self.n.dbs["db-rec"][0]
        self.assertEqual(fila["Próxima fecha"], hoy.replace(day=F.dia_del_mes(23, hoy)).isoformat())
        F.marcar_hecho(self.n, F.recordatorios(self.n, self.bot.bases)[0])
        siguiente = F.sumar_meses(hoy.replace(day=1), 1)
        self.assertEqual(fila["Próxima fecha"], siguiente.replace(day=F.dia_del_mes(23, siguiente)).isoformat())
        self.assertEqual(F.proxima_fecha(31, "", F.date(2027, 2, 10)), F.date(2027, 2, 28))

    def test_proyeccion(self):
        F.fijar_tipo_de_cambio("RUB", 0.05)
        self.rec("Apartamento", 28, "Pago", 2000, "PEN", "Vivienda")
        self.di("/limite 20")
        self.di("+3000 sueldo")
        self.toca("No esta vez")
        self.di("100 supermercado")
        p = F.proyeccion(self.n, self.bot.bases)
        dias = (F.mes().hasta - F.hoy()).days
        self.assertEqual(p["cierre"], round(3000 - 100 - 2000 - 20 * dias, 2))
        self.assertIn("cierras el mes con", self.di("/proyeccion"))
        self.assertIn("cierras el mes con", self.di("/start"))

    def test_plan_de_deudas(self):
        lista = [{"deuda": "A", "saldo_s": 1000, "cuota": 100, "moneda": "PEN", "tasa": 0},
                 {"deuda": "B", "saldo_s": 300, "cuota": 0, "moneda": "PEN", "tasa": 0}]
        base = F.plan_deudas(lista, 0)
        self.assertEqual((base["meses"], base["sin_pagar"]), (13, []))   # la cuota de A pasa a B al terminar
        self.assertEqual(F.plan_deudas([lista[1]], 0)["sin_pagar"], ["B"])   # sin cuota ni extra no termina
        p = F.plan_deudas(lista, 100)
        self.assertEqual(p["fin"]["B"], 3)                      # el extra va primero a la mas chica
        self.assertEqual(p["meses"], 7)                         # 1300 / 200 al mes
        con_tasa = F.plan_deudas([{"deuda": "C", "saldo_s": 1200, "cuota": 100, "moneda": "PEN", "tasa": 0.12}], 0)
        self.assertGreater(con_tasa["meses"], 12)
        self.assertGreater(con_tasa["intereses"], 0)
        self.di("/deuda Banco SIP 16224")
        self.n.dbs["db-deu"][0]["Cuota mensual"] = 1500
        t = self.di("/plan 300 usd")
        self.assertIn("Solo con las cuotas", t)
        self.assertIn("Con $ 300.00 extra al mes", t)

    def test_ingresos_por_fuente_y_grafico(self):
        self.di("+1000 usd facebook")
        self.di("+500 freshtunes")
        t = self.di("/ingresos")
        self.assertIn("Facebook", t)
        self.assertIn("Freshtunes", t)
        self.assertIn("Por mes", t)
        self.assertIn("todavía no hay gastos", self.di("/grafico"))
        self.di("120 supermercado")
        self.di("/grafico")
        self.assertTrue(self.tg.fotos[-1][1].startswith("https://quickchart.io/chart?"))

    def test_alerta_de_tipo_de_cambio(self):
        a = F.ajustes()
        h = {}
        for k in range(15):
            h["2026-09-%02d" % (k + 1)] = {"USD": 3.4, "EUR": 3.9, "RUB": 3.4 / 80}
        h["2026-09-20"] = {"USD": 3.4, "EUR": 3.9, "RUB": 3.4 / 86}      # 86 rublos por dolar: +7.5%
        a["historial_tc"] = h
        F._guardar_ajustes(a)
        self.assertIn("El dólar está alto en rublos", F.alerta_cambio())
        h["2026-09-20"] = {"USD": 3.4, "EUR": 3.9, "RUB": 3.4 / 80.5}
        a["historial_tc"] = h
        F._guardar_ajustes(a)
        self.assertIsNone(F.alerta_cambio())
